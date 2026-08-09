"""v1.16 operating-finance ledger, migration, returns, and UI tests."""

from __future__ import annotations

import json
import sqlite3

import pytest
from openpyxl import load_workbook

from src.models.connection import connect, transaction
from src.models.finance_queries import (
    get_finance_dashboard,
    get_finance_setup_state,
    get_profit_report,
    list_counterparty_balances,
    list_finance_categories,
    list_financial_accounts,
    list_operating_entries,
)
from src.models.migrations import migrate_database
from src.models.queries import get_receivables
from src.models.schema import SCHEMA_VERSION
from src.services.exceptions import DataConflictError, ValidationError
from src.services.finance_service import FinanceService
from src.services.inventory_service import InventoryService
from src.services.order_service import OrderService
from src.services.payment_service import PaymentService
from src.services.reconciliation_service import ReconciliationService
from src.services.return_service import ReturnService
from src.utils.excel_export import export_finance_to_excel
from src.utils.json_export import export_all_to_json, import_from_json


def _seed(path):
    migrate_database(path)
    with transaction(path) as conn:
        product = conn.execute(
            "INSERT INTO products(series,cpu) VALUES ('ThinkPad','i7')"
        ).lastrowid
        customer = conn.execute(
            "INSERT INTO customers(name) VALUES ('客户A')"
        ).lastrowid
        supplier = conn.execute(
            "INSERT INTO suppliers(name) VALUES ('供应商A')"
        ).lastrowid
    return product, customer, supplier


def _enabled(path, opening=0):
    return FinanceService(path).initialize_finance(
        "2026-08-01",
        [{"name": "微信", "opening_balance_cents": opening}],
    )[0]


def _sale(path, *, quantity=2):
    product, customer, supplier = _seed(path)
    account = _enabled(path, 1_000_000)
    batch = InventoryService(path).receive_batch(
        product_id=product,
        purchase_price_cents=300_000,
        quantity=quantity,
        date="2026-08-01",
        supplier_id=supplier,
    )
    quote = OrderService(path).create_quote(
        batch_id=batch,
        customer_id=customer,
        quote_price_cents=350_000,
        quote_quantity=quantity,
        quote_date="2026-07-31",
    )
    InventoryService(path).ship_quote(
        quote, shipped_date="2026-08-02"
    )
    return product, customer, supplier, account, batch, quote


def test_fresh_schema_v4_contains_balanced_ledger(tmp_path):
    path = tmp_path / "fresh.db"
    assert migrate_database(path) is None
    conn = connect(path, read_only=True)
    try:
        assert conn.execute("PRAGMA user_version").fetchone()[0] == 4
        assert SCHEMA_VERSION == 4
        tables = {
            row[0]
            for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        assert {
            "ledger_accounts",
            "finance_categories",
            "ledger_entries",
            "ledger_lines",
            "shipment_snapshots",
            "sales_returns",
            "purchase_returns",
        }.issubset(tables)
        assert conn.execute(
            "SELECT COUNT(*) FROM ledger_accounts WHERE is_system=1"
        ).fetchone()[0] == 8
    finally:
        conn.close()


def test_finance_setup_is_required_and_idempotent(tmp_path):
    path = tmp_path / "setup.db"
    product, _, supplier = _seed(path)
    with pytest.raises(ValidationError, match="财务"):
        InventoryService(path).receive_batch(
            product_id=product,
            purchase_price_cents=10_000,
            quantity=1,
            date="2026-08-01",
            supplier_id=supplier,
        )
    account = _enabled(path, 12_345)
    assert list_financial_accounts(db_path=path)[0]["balance_cents"] == 12_345
    with pytest.raises(DataConflictError):
        _enabled(path)
    assert get_finance_setup_state(path)["enabled"]
    assert account > 0


def test_finance_setup_builds_reconciled_opening_balances(tmp_path):
    path = tmp_path / "opening.db"
    product, customer, supplier = _seed(path)
    with transaction(path) as conn:
        batch = conn.execute(
            """
            INSERT INTO batches(
                product_id,purchase_price_cents,quantity,remaining,date,supplier_id
            ) VALUES (?,?,?,?,?,?)
            """,
            (product, 300_000, 2, 1, "2026-07-01", supplier),
        ).lastrowid
        conn.execute(
            "UPDATE suppliers SET balance_cents=600000 WHERE id=?",
            (supplier,),
        )
        conn.execute(
            """
            INSERT INTO quotes(
                batch_id,customer_id,quote_price_cents,quote_quantity,
                received_amount_cents,status,quote_date
            ) VALUES (?,?,?,?,?,?,?)
            """,
            (batch, customer, 350_000, 1, 100_000, "已出库", "2026-07-02"),
        )
    _enabled(path)
    dashboard = get_finance_dashboard(
        "2026-08-01",
        "2026-08-31",
        path,
    )
    assert dashboard["receivable_cents"] == 250_000
    assert dashboard["payable_cents"] == 600_000
    conn = connect(path, read_only=True)
    try:
        assert conn.execute(
            """
            SELECT SUM(l.debit_cents-l.credit_cents)
            FROM ledger_lines l
            JOIN ledger_accounts a ON a.id=l.account_id
            WHERE a.code='INVENTORY'
            """
        ).fetchone()[0] == 300_000
    finally:
        conn.close()


def test_ship_receipt_advance_and_profit_are_not_duplicated(tmp_path):
    path = tmp_path / "sale.db"
    _, customer, supplier, account, _, quote = _sale(path)
    receipt = PaymentService(path).receive_customer_payment(
        customer,
        800_000,
        "2026-08-03",
        account_id=account,
    )
    assert receipt.allocated_cents == 700_000
    assert receipt.unapplied_cents == 100_000
    dashboard = get_finance_dashboard(
        "2026-08-01", "2026-08-31", path
    )
    assert dashboard["sales_cents"] == 700_000
    assert dashboard["cogs_cents"] == 600_000
    assert dashboard["gross_profit_cents"] == 100_000
    assert dashboard["customer_advance_cents"] == 100_000
    assert dashboard["payable_cents"] == 600_000
    conn = connect(path, read_only=True)
    try:
        assert tuple(conn.execute(
            "SELECT shipped_date,unit_cost_cents FROM shipment_snapshots "
            "WHERE quote_id=?",
            (quote,),
        ).fetchone()) == ("2026-08-02", 300_000)
        assert conn.execute(
            "SELECT COUNT(*) FROM ledger_entries "
            "WHERE idempotency_key=?",
            (f"quote:{quote}:ship",),
        ).fetchone()[0] == 1
    finally:
        conn.close()
    assert ReconciliationService(path).run().is_clean


def test_customer_and_supplier_advances_auto_apply_fifo(tmp_path):
    path = tmp_path / "advance.db"
    product, customer, supplier = _seed(path)
    account = _enabled(path)
    customer_receipt = PaymentService(path).receive_customer_payment(
        customer, 500_000, "2026-08-01", account_id=account
    )
    assert customer_receipt.unapplied_cents == 500_000
    supplier_payment = PaymentService(path).record_supplier_payment(
        supplier, 400_000, "2026-08-01", account_id=account
    )
    batch = InventoryService(path).receive_batch(
        product_id=product,
        purchase_price_cents=300_000,
        quantity=2,
        date="2026-08-02",
        supplier_id=supplier,
    )
    quote = OrderService(path).create_quote(
        batch_id=batch,
        customer_id=customer,
        quote_price_cents=350_000,
        quote_quantity=2,
        quote_date="2026-08-02",
    )
    InventoryService(path).ship_quote(quote, shipped_date="2026-08-03")
    conn = connect(path, read_only=True)
    try:
        assert conn.execute(
            "SELECT SUM(amount_cents) FROM payment_allocations "
            "WHERE payment_id=?",
            (customer_receipt.payment_id,),
        ).fetchone()[0] == 500_000
        assert conn.execute(
            "SELECT SUM(amount_cents) FROM supplier_payment_allocations "
            "WHERE payment_id=?",
            (supplier_payment,),
        ).fetchone()[0] == 400_000
    finally:
        conn.close()


def test_manual_income_expense_transfer_and_reconcile(tmp_path):
    path = tmp_path / "manual.db"
    _seed(path)
    first = _enabled(path, 100_000)
    second = FinanceService(path).create_account("银行卡")
    categories = {
        row["name"]: row["id"]
        for row in list_finance_categories(db_path=path)
    }
    service = FinanceService(path)
    service.record_income(
        account_id=first,
        category_id=categories["其他收入"],
        amount_cents=12_345,
        entry_date="2026-08-02",
    )
    service.record_expense(
        account_id=first,
        category_id=categories["运费"],
        amount_cents=2_345,
        entry_date="2026-08-02",
    )
    service.transfer(
        from_account_id=first,
        to_account_id=second,
        amount_cents=10_000,
        entry_date="2026-08-03",
    )
    service.reconcile_account(
        account_id=second,
        actual_balance_cents=9_999,
        entry_date="2026-08-03",
        reason="实盘差异",
    )
    dashboard = get_finance_dashboard(
        "2026-08-01", "2026-08-31", path
    )
    assert dashboard["other_income_cents"] == 12_345
    assert dashboard["expense_cents"] == 2_345
    assert dashboard["net_profit_cents"] == 10_000
    assert sum(row["balance_cents"] for row in list_financial_accounts(db_path=path)) == 109_999


def test_full_and_partial_returns_keep_ledger_reconciled(tmp_path):
    path = tmp_path / "returns.db"
    _, customer, supplier, account, batch, quote = _sale(path)
    PaymentService(path).receive_customer_payment(
        customer, 700_000, "2026-08-03", account_id=account
    )
    PaymentService(path).record_supplier_payment(
        supplier, 600_000, "2026-08-03", account_id=account
    )
    ReturnService(path).return_sale(
        quote,
        quantity=1,
        return_date="2026-08-04",
        restock=True,
        cash_refund_cents=350_000,
        refund_account_id=account,
        reason="部分退货",
    )
    ReturnService(path).return_purchase(
        batch,
        quantity=1,
        return_date="2026-08-05",
        cash_refund_cents=300_000,
        refund_account_id=account,
        reason="退回供应商",
    )
    dashboard = get_finance_dashboard(
        "2026-08-01", "2026-08-31", path
    )
    assert dashboard["sales_cents"] == 350_000
    assert dashboard["cogs_cents"] == 300_000
    assert dashboard["gross_profit_cents"] == 50_000
    assert not list_counterparty_balances("customer", path)
    assert not list_counterparty_balances("supplier", path)
    assert ReconciliationService(path).run().is_clean


def test_purchase_cash_refund_cannot_exceed_paid_portion(tmp_path):
    path = tmp_path / "unpaid-return.db"
    product, _, supplier = _seed(path)
    account = _enabled(path)
    batch = InventoryService(path).receive_batch(
        product_id=product,
        purchase_price_cents=300_000,
        quantity=1,
        date="2026-08-01",
        supplier_id=supplier,
    )
    with pytest.raises(ValidationError, match="已付款部分"):
        ReturnService(path).return_purchase(
            batch,
            quantity=1,
            return_date="2026-08-02",
            cash_refund_cents=300_000,
            refund_account_id=account,
            reason="尚未付款",
        )


def test_payment_correction_reverses_cash_and_allocations(tmp_path):
    path = tmp_path / "correct.db"
    _, customer, _, account, _, _ = _sale(path, quantity=1)
    payment = PaymentService(path).receive_customer_payment(
        customer, 100_000, "2026-08-03", account_id=account
    )
    replacement = PaymentService(path).correct_payment(
        payment.payment_id,
        amount_cents=150_000,
        pay_date="2026-08-04",
        account_id=account,
        remark="更正",
        reason="金额录错",
    )
    assert replacement != payment.payment_id
    assert ReconciliationService(path).run().is_clean


def test_v4_json_and_finance_excel_round_trip(tmp_path):
    source = tmp_path / "source.db"
    target = tmp_path / "target.db"
    export = tmp_path / "backup.json"
    workbook = tmp_path / "finance.xlsx"
    _sale(source, quantity=1)
    export_all_to_json(export, source)
    document = json.loads(export.read_text(encoding="utf-8"))
    assert document["schema_version"] == 4
    assert "ledger_entries" in document["data"]
    migrate_database(target)
    ok, _, _ = import_from_json(export, target)
    assert ok
    assert get_finance_dashboard(
        "2026-08-01", "2026-08-31", target
    ) == get_finance_dashboard("2026-08-01", "2026-08-31", source)
    export_finance_to_excel(
        workbook,
        date_from="2026-08-01",
        date_to="2026-08-31",
        db_path=target,
    )
    assert workbook.exists() and workbook.stat().st_size > 0
    book = load_workbook(workbook, read_only=True)
    assert {
        "账户流水",
        "日常收支",
        "客户往来",
        "供应商往来",
        "利润明细",
        "审计摘要",
    }.issubset(book.sheetnames)
    book.close()


def test_audit_filters_and_replacement_status(tmp_path):
    path = tmp_path / "audit-filter.db"
    _, customer, _, account, _, _ = _sale(path, quantity=1)
    original = PaymentService(path).receive_customer_payment(
        customer,
        100_000,
        "2026-08-03",
        account_id=account,
    ).payment_id
    replacement = PaymentService(path).correct_payment(
        original,
        amount_cents=150_000,
        pay_date="2026-08-04",
        account_id=account,
        remark="更正",
        reason="金额录错",
    )
    rows = list_operating_entries(
        event_type="customer_receipt",
        status="corrected",
        source_id=replacement,
        db_path=path,
    )
    assert len(rows) == 1
    assert rows[0]["source_id"] == str(replacement)


def test_finance_tab_has_five_pages(qapp, tmp_path):
    from src.ui.finance_tab import FinanceTab

    path = tmp_path / "ui.db"
    _seed(path)
    _enabled(path)
    tab = FinanceTab(db_path=path)
    tab.refresh()
    assert tab.section_tabs.count() == 5
    assert [
        tab.section_tabs.tabText(index) for index in range(5)
    ] == ["经营总览", "日常收支", "应收应付", "利润分析", "流水审计"]


def test_customer_receivable_after_shipment(tmp_path):
    """出库后客户应收应正确反映在 get_receivables() 中"""
    db = tmp_path / "test.db"
    product, customer, supplier = _seed(db)
    account = _enabled(db)

    inv = InventoryService(db)
    batch_id = inv.receive_batch(
        product_id=product,
        purchase_price_cents=500000,
        quantity=10,
        date="2026-08-02",
        settlement_mode="paid",
        account_id=account,
    )
    order = OrderService(db)
    quote_id = order.create_quote(
        customer_id=customer,
        batch_id=batch_id,
        quote_price_cents=600000,
        quote_quantity=3,
        quote_date="2026-08-02",
    )
    inv.ship_quote(quote_id, shipped_date="2026-08-03")

    receivables = {r["id"]: r for r in get_receivables(db)}
    assert customer in receivables
    assert receivables[customer]["debt_cents"] == 600000 * 3


def test_customer_receivable_after_payment(tmp_path):
    """收款后客户应收应正确减少"""
    db = tmp_path / "test.db"
    product, customer, supplier = _seed(db)
    account = _enabled(db)

    inv = InventoryService(db)
    batch_id = inv.receive_batch(
        product_id=product,
        purchase_price_cents=500000,
        quantity=10,
        date="2026-08-02",
        settlement_mode="paid",
        account_id=account,
    )
    order = OrderService(db)
    quote_id = order.create_quote(
        customer_id=customer,
        batch_id=batch_id,
        quote_price_cents=600000,
        quote_quantity=3,
        quote_date="2026-08-02",
    )
    inv.ship_quote(quote_id, shipped_date="2026-08-03")

    pay = PaymentService(db)
    pay.receive_customer_payment(
        customer_id=customer,
        amount_cents=1000000,
        pay_date="2026-08-04",
        account_id=account,
    )

    receivables = {r["id"]: r for r in get_receivables(db)}
    assert customer in receivables
    assert receivables[customer]["debt_cents"] == 600000 * 3 - 1000000


def test_customer_balance_cents_unchanged_after_shipment(tmp_path):
    """出库后 customers.balance_cents 应保持为 0"""
    db = tmp_path / "test.db"
    product, customer, supplier = _seed(db)
    account = _enabled(db)

    inv = InventoryService(db)
    batch_id = inv.receive_batch(
        product_id=product,
        purchase_price_cents=500000,
        quantity=10,
        date="2026-08-02",
        settlement_mode="paid",
        account_id=account,
    )
    order = OrderService(db)
    quote_id = order.create_quote(
        customer_id=customer,
        batch_id=batch_id,
        quote_price_cents=600000,
        quote_quantity=3,
        quote_date="2026-08-02",
    )
    inv.ship_quote(quote_id, shipped_date="2026-08-03")

    with transaction(db) as conn:
        row = conn.execute(
            "SELECT balance_cents FROM customers WHERE id=?", (customer,)
        ).fetchone()
        assert row["balance_cents"] == 0


def test_customer_balance_cents_unchanged_after_payment(tmp_path):
    """收款后 customers.balance_cents 应保持为 0"""
    db = tmp_path / "test.db"
    product, customer, supplier = _seed(db)
    account = _enabled(db)

    inv = InventoryService(db)
    batch_id = inv.receive_batch(
        product_id=product,
        purchase_price_cents=500000,
        quantity=10,
        date="2026-08-02",
        settlement_mode="paid",
        account_id=account,
    )
    order = OrderService(db)
    quote_id = order.create_quote(
        customer_id=customer,
        batch_id=batch_id,
        quote_price_cents=600000,
        quote_quantity=3,
        quote_date="2026-08-02",
    )
    inv.ship_quote(quote_id, shipped_date="2026-08-03")

    pay = PaymentService(db)
    pay.receive_customer_payment(
        customer_id=customer,
        amount_cents=1000000,
        pay_date="2026-08-04",
        account_id=account,
    )

    with transaction(db) as conn:
        row = conn.execute(
            "SELECT balance_cents FROM customers WHERE id=?", (customer,)
        ).fetchone()
        assert row["balance_cents"] == 0
